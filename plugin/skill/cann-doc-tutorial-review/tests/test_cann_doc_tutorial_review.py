"""cann-doc-tutorial-review 脚本单测(纯逻辑;codecheck 用系统 grep)。"""
import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import _state          # noqa: E402
import find_tutorials  # noqa: E402
import codecheck       # noqa: E402
import render_report   # noqa: E402


def _redirect(monkeypatch, tmp_path):
    monkeypatch.setattr(_state, "DOCCHECK_ROOT", tmp_path / "doccheck")


# ---- find_tutorials:命中开发指南、排除 quickstart/README/参考、跳噪声目录 ----

def test_find_tutorials(tmp_path):
    d = tmp_path / "docs" / "zh"
    (d / "develop").mkdir(parents=True)
    (d / "develop" / "aicore_develop_guide.md").write_text("# AI Core算子开发指南\n", encoding="utf-8")  # 文件名命中
    (d / "advanced.md").write_text("## 进阶\n内容", encoding="utf-8")                                    # 标题命中
    (tmp_path / "QUICKSTART.md").write_text("# 快速入门", encoding="utf-8")                               # 排除
    (d / "op_list.md").write_text("# 算子清单", encoding="utf-8")                                         # 排除(参考清单)
    (tmp_path / "README.md").write_text("# Develop Guide here", encoding="utf-8")                          # 文件名排除
    (tmp_path / "build").mkdir()
    (tmp_path / "build" / "x_develop_guide.md").write_text("# guide", encoding="utf-8")                   # 噪声目录,跳

    paths = [x["path"] for x in find_tutorials.find_tutorials(str(tmp_path))]
    assert "docs/zh/develop/aicore_develop_guide.md" in paths
    assert "docs/zh/advanced.md" in paths
    assert "QUICKSTART.md" not in paths
    assert "docs/zh/op_list.md" not in paths
    assert "README.md" not in paths
    assert all("build/" not in p for p in paths)


# ---- discover_docs(范围默认):限定 docs/ + 取全部技术文档(非教程启发式)----

def test_discover_docs_scope(tmp_path):
    (tmp_path / "docs" / "zh" / "context").mkdir(parents=True)
    (tmp_path / "docs" / "zh" / "context" / "数据类型.md").write_text("# 数据类型\n", encoding="utf-8")  # 非教程,也要取
    (tmp_path / "docs" / "QUICKSTART.md").write_text("# 快速入门\n", encoding="utf-8")                     # docs/ 下,取
    (tmp_path / "docs" / "build").mkdir()
    (tmp_path / "docs" / "build" / "x.md").write_text("# x", encoding="utf-8")                              # 噪声目录,跳
    (tmp_path / "myop").mkdir()
    (tmp_path / "myop" / "myop_develop_guide.md").write_text("# 算子开发指南\n", encoding="utf-8")         # 算子目录里的教程,不取

    paths = [x["path"] for x in find_tutorials.discover_docs(str(tmp_path), "docs")]
    assert "docs/zh/context/数据类型.md" in paths        # docs/ 下非教程也取
    assert "docs/QUICKSTART.md" in paths                  # docs/ 全量
    assert all("build/" not in p for p in paths)          # 噪声目录跳
    assert all(p.startswith("docs/") for p in paths)      # 严格限定 docs/ 子树
    assert "myop/myop_develop_guide.md" not in paths      # 算子目录的教程,不越界取


# ---- codecheck triage:占位符 / 外部命令 / 仓内对象 ----

def test_triage():
    assert codecheck.triage("<repo>", "path") == "placeholder"
    assert codecheck.triage("$ASCEND_HOME/x", "path") == "placeholder"
    assert codecheck.triage("cmake ..", "command") == "external"
    assert codecheck.triage("pip install x", "command") == "external"
    assert codecheck.triage("build.sh --pkg", "command") == "repo_object"
    assert codecheck.triage("build/foo.run", "path") == "generated"


# ---- codecheck 取证:强命中→CONSISTENT;缺失+近似变体→SUSPECTED 带变体 ----

def test_codecheck_strong_and_variant(tmp_path):
    (tmp_path / "build.sh").write_text(
        'case "$1" in\n  --opkernel) echo ok ;;\nesac\n', encoding="utf-8")
    # 文档写的 --opkernel 真有(case 分支=强证据)
    r1 = codecheck.check(str(tmp_path), "--opkernel", "flag")
    assert r1["triage"] == "repo_object"
    assert r1["grade"] == "strong" and r1["suggested_verdict"] == "CONSISTENT"
    # 文档写的 --opkernel_test 没有,但近似变体 --opkernel 在 → SUSPECTED + 变体(agent 可升 CONFIRMED)
    r2 = codecheck.check(str(tmp_path), "--opkernel_test", "flag")
    assert r2["suggested_verdict"] == "SUSPECTED"
    assert "--opkernel" in r2["near_variants"]
    # codecheck 绝不自己下 CONFIRMED
    assert r2["suggested_verdict"] != "CONFIRMED_MISMATCH"


def test_codecheck_path_missing_vs_present(tmp_path):
    (tmp_path / "real.txt").write_text("x", encoding="utf-8")
    assert codecheck.check(str(tmp_path), "real.txt", "path")["suggested_verdict"] == "CONSISTENT"
    assert codecheck.check(str(tmp_path), "docs/missing.yml", "path")["suggested_verdict"] == "SUSPECTED"


# ---- 自校验闸:可量化无代码位置 / 不可量化无判例·steelman / 确认讲错无外部反证 → 不过 ----

def test_self_check_gate():
    ok_q = {"cls": "quantifiable", "axis": "trustworthy", "quote": "x", "verdict": "CONFIRMED_MISMATCH",
            "improvement": "改", "code_location": "build.sh:3"}
    assert _state.self_check_finding(ok_q) == []
    bad_q = {**ok_q, "code_location": None}
    assert any("code_location" in p for p in _state.self_check_finding(bad_q))   # 可量化缺陷必带代码位置

    ok_n = {"cls": "non_quantifiable", "axis": "learnable", "quote": "x", "verdict": "TEACHING_JUDGMENT",
            "improvement": "改", "precedent": "读者卡在…", "steelman": "已反论"}
    assert _state.self_check_finding(ok_n) == []
    assert any("precedent" in p for p in _state.self_check_finding({**ok_n, "precedent": None}))
    assert any("steelman" in p for p in _state.self_check_finding({**ok_n, "steelman": None}))
    # 确认讲错必带外部反证,否则违规
    cw = {**ok_n, "verdict": "CONFIRMED_CONCEPT_WRONG"}
    assert any("external_evidence" in p for p in _state.self_check_finding(cw))


# ---- render:自校验闸把无证据的踢进「待补」;两段分开;轴档 ----

def test_render_two_sections_and_gate(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "docs/zh/develop/g.md", "type": "教学型", "audience": "已读 quickstart",
                               "code_root": "/x"})
    # 1 可量化确认(过闸) 1 教学判断(过闸) 1 无证据(踢待补)
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错",
                                 "source": "code_mismatch", "quote": "--opkernel_test", "verdict": "CONFIRMED_MISMATCH",
                                 "evidence_grade": "strong", "code_location": "build.sh:3", "improvement": "改成 --opkernel"})
    _state.add_finding("ops-x", {"cls": "non_quantifiable", "axis": "learnable", "form": "缺",
                                 "quote": "适当配置 tiling", "verdict": "TEACHING_JUDGMENT",
                                 "precedent": "读者不知 tiling 取值依据", "steelman": "上文未给公式",
                                 "improvement": "补 tiling 计算约束"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "quote": "无证据条",
                                 "verdict": "CONFIRMED_MISMATCH", "improvement": "x"})   # 缺 code_location → 待补
    md = render_report.render("ops-x")
    assert "对不上代码" in md and "讲得不到位" in md    # 事实问题走大类表格,教学判断走大类详列
    assert "--opkernel_test" in md and "适当配置 tiling" in md
    assert "待补" in md and "无证据条" in md          # 无证据的被踢进待补,不进事实问题正文
    assert "不合格" in md                             # trustworthy 轴有 CONFIRMED → 不合格


# ---- TE-1:axes_evaluated 缺省=全评(向后兼容);给了子集 → 未列轴标「本轮未评」 ----

def test_unevaluated_axis(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "docs/zh/develop/g.md", "type": "教学型", "audience": "x",
                               "code_root": "/x", "axes_evaluated": ["trustworthy"]})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错",
                                 "source": "code_mismatch", "quote": "x", "verdict": "SUSPECTED",
                                 "evidence_grade": "medium", "code_location": "build.sh:1", "improvement": "改"})
    md = render_report.render("ops-x")
    assert md.count("本轮未评") == 4                 # 5 轴里只评了 trustworthy,其余 4 轴未评(MD 总评表)
    # 缺省字段 → 向后兼容(不出现「本轮未评」)
    _state.save_meta("ops-x", {"doc": "docs/zh/develop/g.md", "type": "教学型", "audience": "x", "code_root": "/x"})
    assert "本轮未评" not in render_report.render("ops-x")


# ---- TE-2:同一教程行号、同 cls 的跨轴重复 → 折叠保留证据最强,记 also_hit ----

def test_dedup_by_docline(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "docs/zh/develop/guide.md", "type": "教学型", "audience": "x", "code_root": "/x"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错", "source": "code_mismatch",
                                 "quote": "信得过版", "verdict": "CONFIRMED_MISMATCH", "evidence_grade": "strong",
                                 "code_location": "ops-nn/docs/zh/develop/guide.md:100 真证据", "improvement": "改"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "readable", "form": "错", "source": "self_contradiction",
                                 "quote": "读得懂版", "verdict": "SUSPECTED", "evidence_grade": "medium",
                                 "code_location": "guide.md:100 弱证据", "improvement": "改2"})
    md = render_report.render("ops-x")
    assert "另命中此处的轴:读得懂" in md             # readable 被折叠进 trustworthy,留痕
    assert "| 信得过 | 不合格 | 1 |" in md            # 证据最强(CONFIRMED/trustworthy)胜出
    assert "| 读得懂 | 合格 | 0 |" in md             # readable 那条已折叠走,不再计数
    assert "信得过版" in md and "读得懂版" not in md   # 正文只留胜出条的原文
    # 不同行不折叠
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "readable", "form": "错", "source": "code_mismatch",
                                 "quote": "另一行", "verdict": "SUSPECTED", "evidence_grade": "weak",
                                 "code_location": "guide.md:200", "improvement": "改3"})
    assert "另一行" in render_report.render("ops-x")


# ---- TE-3:render_html 出卡片/折叠/三态色标,且自校验闸把无证据条落「待补」段 ----

def test_render_html(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "docs/zh/develop/g.md", "type": "教学型", "audience": "x", "code_root": "/x"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错", "source": "code_mismatch",
                                 "quote": "GlobalTensor<T> x_;", "verdict": "CONFIRMED_MISMATCH", "evidence_grade": "strong",
                                 "code_location": "g.md:5", "improvement": "改", "impact": "blocker", "category": "C4.1"})
    h = render_report.render_html("ops-x")
    assert h.strip().endswith("</html>")
    assert "var DATA =" in h and "算子文档体检报告" in h         # 设计引擎(按问题类型)注入
    assert '"cat": "C3"' in h                                   # category C4.1 → 8 类 C3(代码片段)
    assert "docs/zh/develop/g.md" in h                          # doc 进 DATA
    assert "GlobalTensor<T>" in h                               # 原文进 DATA(引擎 JS 运行时再 esc)
    assert '"impact": "blocker"' in h and '"suspected": false' in h


# ---- linkcheck(T0):死文件链 + 死锚点 ----

def test_linkcheck(tmp_path):
    import linkcheck
    (tmp_path / "a.md").write_text(
        "[死链](./nope.cpp)\n[活锚](./b.md#标题)\n[死锚](./b.md#不存在)\n", encoding="utf-8")
    (tmp_path / "b.md").write_text("## 标题\n正文", encoding="utf-8")
    fs = linkcheck.find_broken_links(str(tmp_path))
    assert len(fs) == 2 and {f["category"] for f in fs} == {"C1"}   # 一条死文件 + 一条死锚(均 C1),活锚不报
    assert sum("锚点不存在" in f["code_location"] for f in fs) == 1   # 恰一条死锚
    assert all(f["impact"] == "misleading" for f in fs)


# ---- support_table_check(T0):同算子两篇文档支持表自相矛盾 ----

def test_support_table_contradiction(tmp_path):
    import support_table_check
    op = tmp_path / "myop"
    (op / "op_host").mkdir(parents=True)
    (op / "README.md").write_text("| <term>Atlas A2 训练系列产品</term> | √ |\n", encoding="utf-8")
    (op / "docs").mkdir()
    (op / "docs" / "aclnnMyOp.md").write_text("| <term>Atlas A2 训练系列产品</term> | × |\n", encoding="utf-8")
    fs = support_table_check.check(str(tmp_path))
    assert len(fs) == 1
    assert fs[0]["category"] == "C5" and fs[0]["verdict"] == "CONFIRMED_MISMATCH"


# ---- T0 脚本 --under 跟随 P0 范围(TG-2 修复)----

def test_t0_under_scope(tmp_path):
    import linkcheck, support_table_check
    (tmp_path / "docs").mkdir()
    (tmp_path / "docs" / "g.md").write_text("[死链](./none.cpp)\n", encoding="utf-8")
    op = tmp_path / "myop"
    (op / "op_host").mkdir(parents=True)
    (op / "x.md").write_text("[死链](./none.cpp)\n", encoding="utf-8")
    # linkcheck:--under docs 只扫 docs/(1 条);不限范围两条都报
    assert len(linkcheck.find_broken_links(str(tmp_path), under="docs")) == 1
    assert len(linkcheck.find_broken_links(str(tmp_path))) == 2
    # support_table_check:--under docs 限定中央文档(无产品支持表)→ 0
    assert support_table_check.check(str(tmp_path), under="docs") == []


# ---- render:阻断排在误导前 + 影响列存在 ----

def test_render_impact_sort(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("r", {"doc": "t.md", "axes_evaluated": ["trustworthy"]})
    base = dict(cls="quantifiable", axis="trustworthy", verdict="CONFIRMED_MISMATCH",
                improvement="fix", code_location="x:1")
    _state.save_findings("r", [
        {**base, "quote": "MISLEAD_ONE", "impact": "misleading"},
        {**base, "quote": "BLOCKER_ONE", "impact": "blocker"},
    ])
    md = render_report.render("r")
    assert "🔴 阻断" in md and "严重度" in md
    assert md.index("BLOCKER_ONE") < md.index("MISLEAD_ONE")   # 阻断排前


def test_drop_minor_by_default(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("r", {"doc": "t.md", "axes_evaluated": ["trustworthy"]})
    base = dict(cls="quantifiable", axis="trustworthy", verdict="CONFIRMED_MISMATCH",
                improvement="fix", code_location="x:1")
    _state.save_findings("r", [
        {**base, "quote": "MINOR_DROP", "impact": "minor"},
        {**base, "quote": "BLOCKER_KEEP", "impact": "blocker"},
    ])
    assert render_report.WITH_MINOR is False                       # 默认舍弃
    md = render_report.render("r")
    assert "BLOCKER_KEEP" in md and "MINOR_DROP" not in md         # 默认丢 minor
    monkeypatch.setattr(render_report, "WITH_MINOR", True)
    assert "MINOR_DROP" in render_report.render("r")               # --with-minor 可保留


# ── 分类迁移:旧 8 类码 → 细类(S1–S7);锁死映射,防回归 ──
def test_design_sub_migrates_legacy_codes():
    from scripts.render_report import _design_sub
    expect = {"C1": "S3", "C2": "S1", "C3": "S2", "C4": "S4",
              "C5": "S1", "C6": "S2", "J1": "S5", "J2": "S6"}
    for old, new in expect.items():
        assert _design_sub({"category": old}) == new, f"{old} 应映射到 {new}"


# ── 7 类错误类型模型:S6(描述不清晰)/S7(关键信息缺失)由 finder 直出或旧教学条启发拆分 ──
def test_taxonomy_seven_error_types():
    assert set(_state.SUB_ORDER) == {"S1", "S2", "S3", "S4", "S5", "S6", "S7"}
    assert len(_state.ERR_TYPE_LABEL) == 7
    assert _state.sub_group("S7") == "教学"
    assert _state.err_type("S1") == "文档接口声明与代码不符"
    assert _state.err_type("S4") == "同源文档内容冲突"
    assert _state.err_type("S7") == "关键信息要素缺失"
    # 错误类型全称 = 细类名(报告/HTML/Excel 同一套词,防漂移)
    for sub in _state.SUB_ORDER:
        assert _state.ERR_TYPE_LABEL[sub] == _state.SUBS[sub][0]


def test_design_sub_accepts_explicit_s6_s7():
    from scripts.render_report import _design_sub
    assert _design_sub({"sub": "S6"}) == "S6"
    assert _design_sub({"sub": "S7"}) == "S7"


def test_legacy_teaching_j2_split():
    """旧纯 category 教学条(J2)按缺失信号拆 S6/S7;显式 sub 不被启发改写。"""
    from scripts.render_report import _design_sub
    assert _design_sub({"category": "J2", "prob": "缺前置步骤说明"}) == "S7"   # 缺要素 → S7
    assert _design_sub({"category": "J2", "prob": "未写明需进入哪个目录"}) == "S7"
    assert _design_sub({"category": "J2", "prob": "措辞含糊,读者难判断"}) == "S6"   # 说不清 → S6
    assert _design_sub({"category": "J2"}) == "S6"                                  # 无信号默认 S6
    assert _design_sub({"category": "J2", "sub": "S6", "prob": "缺前置步骤"}) == "S6"  # 直出不拆


def test_design_sub_prefers_explicit_sub():
    """finder 直出 sub 时不走兜底猜测。"""
    from scripts.render_report import _design_sub
    assert _design_sub({"sub": "S5", "category": "C1"}) == "S5"


def test_design_sub_legacy_leaf_codes():
    """老叶子码(C7–C9/J3、带小数点)先归 8 类再落细类(S1–S7)。"""
    from scripts.render_report import _design_sub
    assert _design_sub({"category": "C7"}) == "S1"   # 旧支持表 → 契约类
    assert _design_sub({"category": "C9.2"}) == "S2"  # 旧数值/行为 → 可执行类
    assert _design_sub({"category": "J3"}) == "S6"


def test_html_carries_sub_not_legacy_type(tmp_path, monkeypatch):
    """HTML DATA 必须带 sub;engine 已不消费 type/check。"""
    import json, re
    from scripts import render_report
    _redirect(monkeypatch, tmp_path)          # 必须重定向,否则产物会写进真实仓库
    repo = "subsync"
    _state.save_meta(repo, {"doc": "d.md", "code_root": "/r"})
    _state.save_findings(repo, [{
        "idx": 1, "cls": "quantifiable", "axis": "trustworthy", "category": "C5",
        "quote": "支持表 310P √", "verdict": "CONFIRMED_MISMATCH",
        "code_location": "op_def.cc:5", "improvement": "去掉勾",
        "impact": "blocker",
    }])
    html = render_report.render_html(repo)
    data = json.loads(re.search(r"var DATA = (\[.*?\]);", html, re.S).group(1))
    assert data[0]["sub"] == "S1", "C5(支持表) 应落到 S1 契约类"


# ── 防漂移:Python 侧分类模型(_state) 必须与 HTML 引擎里的 JS GROUPS/SUBS 完全一致 ──
def test_taxonomy_python_matches_html_engine():
    """分类模型存在两处(Python 供 MD/映射,JS 供自包含 HTML)。此测试锁死二者一致,
    防止改了一处忘另一处 → MD 与 HTML 说不同分类(本次重构要根除的正是这个问题)。"""
    import re
    from pathlib import Path
    from scripts import _state
    tpl = (Path(__file__).resolve().parent.parent / "templates" / "report-engine.html").read_text(encoding="utf-8")

    js_subs = dict(re.findall(r"(S\d):\['([^']+)','([^']+)'\]", tpl.replace(" ", "")) and
                   [(m[0], (m[1], m[2])) for m in re.findall(r"(S\d):\['([^']+)','([^']+)'\]", tpl.replace(" ", ""))])
    assert js_subs, "未能从模板解析出 JS SUBS"
    assert js_subs == _state.SUBS, f"细类不一致\nJS : {js_subs}\nPY : {_state.SUBS}"

    js_groups = {m[0]: (m[1], m[2], m[3]) for m in
                 re.findall(r"'(代码|文档|教学)':\['([^']+)','([^']+)','([^']+)'\]", tpl.replace(" ", ""))}
    assert js_groups == _state.GROUPS, f"大类不一致\nJS : {js_groups}\nPY : {_state.GROUPS}"

    js_order = re.search(r"varSUBORDER=\[([^\]]+)\]", tpl.replace(" ", "")).group(1).replace("'", "").split(",")
    assert js_order == _state.SUB_ORDER


# ---- 确定度单一真相源:verdict 说了算,老 conf 字段不得反悔 ----

def test_conf_field_cannot_contradict_verdict(tmp_path, monkeypatch):
    """老产物里 verdict=SUSPECTED 却 conf=true 时,报告必须按 verdict 判「疑似」。"""
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "d.md", "type": "教学型", "audience": "x", "code_root": "/x"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错",
                                 "source": "code_mismatch", "quote": "疑似条", "verdict": "SUSPECTED",
                                 "evidence_grade": "weak", "code_location": "a.sh:1",
                                 "improvement": "查证", "impact": "misleading", "conf": True})
    assert render_report._conf_label({"verdict": "SUSPECTED", "conf": True}) == "疑似"
    h = render_report.render_html("ops-x")
    assert '"suspected": true' in h


def test_confirmed_verdict_marks_confirmed(tmp_path, monkeypatch):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "d.md", "type": "教学型", "audience": "x", "code_root": "/x"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错",
                                 "source": "code_mismatch", "quote": "确认条",
                                 "verdict": "CONFIRMED_MISMATCH", "evidence_grade": "strong",
                                 "code_location": "a.sh:1", "improvement": "改", "impact": "blocker",
                                 "conf": False})     # 老字段说「疑似」,verdict 说「确认」→ 以 verdict 为准
    assert '"suspected": false' in render_report.render_html("ops-x")


def test_fig_and_conseq_are_optional(tmp_path, monkeypatch):
    """省略 fig / conseq 不该让渲染失败,后果句走 impact 默认。"""
    _redirect(monkeypatch, tmp_path)
    _state.save_meta("ops-x", {"doc": "d.md", "type": "教学型", "audience": "x", "code_root": "/x"})
    _state.add_finding("ops-x", {"cls": "quantifiable", "axis": "trustworthy", "form": "错",
                                 "source": "code_mismatch", "quote": "无 fig 条",
                                 "verdict": "CONFIRMED_MISMATCH", "evidence_grade": "strong",
                                 "code_location": "a.sh:1", "improvement": "改", "impact": "blocker"})
    h = render_report.render_html("ops-x")
    assert '"fig"' not in h                 # 没给就不注入,引擎自己不画
    assert '"conseq": ""' not in h          # 走 impact 默认句,不是空串


# ============================ linkcheck:本地克隆护栏(越出仓根/submodule/LFS 不误报) ============================

def test_linkcheck_submodule_lfs_guards(tmp_path):
    """submodule 未 init / LFS 未拉取的缺失目标不该当死链报;真死链仍报。"""
    import linkcheck
    docs = tmp_path / "docs"
    docs.mkdir(parents=True)
    (tmp_path / ".gitmodules").write_text(
        '[submodule "third_party/nn"]\n\tpath = third_party/nn\n', encoding="utf-8")
    (tmp_path / ".gitattributes").write_text("*.png filter=lfs diff=lfs merge=lfs -text\n", encoding="utf-8")
    (docs / "a.md").write_text(
        "[子模块](../third_party/nn/x.md)\n"      # 命中 submodule 前缀 → 跳过
        "[LFS图](./fig.png)\n"                     # 命中 *.png lfs → 跳过
        "[真死链](./nope.md)\n", encoding="utf-8")  # 仓内缺失 → 报
    fs = linkcheck.find_broken_links(str(tmp_path), under="docs")
    assert len(fs) == 1 and fs[0]["impact"] == "misleading"
    assert "nope.md" in fs[0]["code_location"]


def test_linkcheck_out_of_tree_downgrade(tmp_path):
    """越出仓根的引用(跨仓/上级文档集)本地快照无法证实 → 降 SUSPECTED/minor,不算「确认死链」。"""
    import linkcheck
    (tmp_path / "docs").mkdir(parents=True)
    (tmp_path / "docs" / "a.md").write_text(
        "[跨仓](../../sibling_doc.md)\n"
        "[仓内缺失](./nope.md)\n", encoding="utf-8")
    fs = linkcheck.find_broken_links(str(tmp_path), under="docs")
    outs = [f for f in fs if "越出仓根" in f["code_location"]]
    nope = [f for f in fs if "nope.md" in f["code_location"]]
    assert outs and outs[0]["verdict"] == "SUSPECTED" and outs[0]["impact"] == "minor"
    assert nope and nope[0]["verdict"] == "CONFIRMED_MISMATCH" and nope[0]["impact"] == "misleading"


# ============================ 导出 Excel:明细 7 列 / 错误类型全称 / 降级 CSV+SVG ============================

def _seed_three_sub_repo(tmp_path, monkeypatch, repo="ops-x"):
    _redirect(monkeypatch, tmp_path)
    _state.save_meta(repo, {"doc": "docs/zh/g.md", "type": "教学型", "audience": "x",
                            "code_root": "/x", "axes_evaluated": ["trustworthy", "learnable"]})
    _state.add_finding(repo, {"cls": "quantifiable", "axis": "trustworthy", "form": "错",
                              "source": "code_mismatch", "quote": "--flag 不存在", "sub": "S1",
                              "verdict": "CONFIRMED_MISMATCH", "evidence_grade": "strong",
                              "code_location": "build.sh:3", "improvement": "build.sh 无该选项,应删或补实现",
                              "fix": "从文档删除 --flag", "impact": "misleading", "doc": "docs/zh/g.md"})
    _state.add_finding(repo, {"cls": "non_quantifiable", "axis": "learnable", "form": "缺",
                              "quote": "默认分 8", "sub": "S7", "verdict": "TEACHING_JUDGMENT",
                              "precedent": "读者不知工作目录", "steelman": "日志路径可反推",
                              "improvement": "补 cd 步骤", "fix": "补 cd 步骤", "impact": "misleading",
                              "doc": "docs/zh/g.md"})
    _state.add_finding(repo, {"cls": "non_quantifiable", "axis": "learnable", "form": "糊",
                              "quote": "按需配置", "sub": "S6", "verdict": "TEACHING_JUDGMENT",
                              "precedent": "读者不知取值", "steelman": "上文未给依据",
                              "improvement": "写明约束", "impact": "misleading", "doc": "docs/zh/g.md"})
    return repo


def test_export_build_rows_and_excel(tmp_path, monkeypatch):
    import importlib.util
    from scripts import export_excel
    repo = _seed_three_sub_repo(tmp_path, monkeypatch)
    rows = export_excel.build_rows(repo)
    assert rows[0][:4] == ["ops-x", "docs/zh/g.md", "对不上代码", "文档接口声明与代码不符"]
    et_col = export_excel.HEADERS.index("错误类型")
    labels = {r[et_col] for r in rows}
    assert {"文档接口声明与代码不符", "文档内容描述不清晰", "关键信息要素缺失"} <= labels
    # 教学缺要素条应落到 S7 全称(而非 S6)
    assert "关键信息要素缺失" in labels

    produced = export_excel.export(repo)
    if importlib.util.find_spec("openpyxl"):
        xlsx = [p for p in produced if p.suffix == ".xlsx"]
        assert xlsx, "openpyxl 在时应产出 .xlsx"
        import openpyxl
        wb = openpyxl.load_workbook(xlsx[0], read_only=True)
        assert list(wb["问题明细"].iter_rows(max_row=1, values_only=True))[0] == tuple(export_excel.HEADERS)
        vals = list(wb["问题明细"].iter_rows(values_only=True))
        assert any(r[3] == "文档接口声明与代码不符" for r in vals[1:])
    else:
        assert any(p.suffix == ".csv" for p in produced), "无 openpyxl 时降级 csv"


def test_export_fallback_csv_and_svg(tmp_path):
    """纯 stdlib 降级:CSV 明细含表头,SVG 图可写。"""
    from scripts import export_excel
    out = tmp_path / "ops-x"
    out.mkdir(parents=True)
    rows = [["ops-x", "docs/zh/g.md", "对不上代码", "文档接口声明与代码不符",
             "接口对不上", "原文 x\n证据 y", "改 z"]]
    st = {"total": 1, "files": ["docs/zh/g.md"], "sev": {"阻断": 0, "误导": 1},
          "conf": {"确认": 1, "疑似": 0},
          "by_sub": {"文档接口声明与代码不符": {"n": 1, "blocker": 0, "misleading": 1}},
          "by_group": {"对不上代码": {"n": 1, "blocker": 0}}}
    produced = export_excel.write_csv_fallback(out / "ops-x.xlsx", "ops-x", rows, st)
    assert all(p.is_file() for p in produced)
    assert produced[0].read_text(encoding="utf-8-sig").splitlines()[0].startswith("仓库名")
    svg = out / "statistics.svg"
    export_excel.write_svg(svg, "统计", ["文档接口声明与代码不符"], [1])
    assert "<svg" in svg.read_text(encoding="utf-8")
