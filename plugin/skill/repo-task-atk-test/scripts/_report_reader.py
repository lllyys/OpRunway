"""读取 ATK xlsx 报告的表格结构。"""

import ast
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


CASE_ID_ALIASES = ("case_id", "编号", "id")
EMPTY_CELLS = {"", "-", "none", "nan", "null"}
TRUE_WORDS = ("true", "pass", "1", "通过")

ACC_PASS = "精度通过"
ACC_DETAIL = "精度详情"
PERF_TIME = "Device性能（us）"
PERF_RATIO = "Device性能比"
PERF_PASS = "device性能通过"
PERF_FLUCT = "性能波动校验结果"
RESULT_COL = "运行结果"
REASON_COL = "失败原因"
RANGE_COL = "输入range"

NODE_METRIC_SUFFIXES = (
    ACC_PASS,
    ACC_DETAIL,
    PERF_TIME,
    PERF_RATIO,
    PERF_PASS,
    PERF_FLUCT,
)
KNOWN_BACKENDS = (
    "pyaclnn",
    "aclnn",
    "npu",
    "gpu",
    "cpu",
    "atb",
    "triton",
    "fusion",
    "dist",
    "kernel",
)
BENCHMARK_SUFFIX = "_benchmark"


def read_workbook(path):
    if load_workbook is None:
        raise RuntimeError("需要 openpyxl：pip install openpyxl")
    return load_workbook(path, data_only=True, read_only=True)


def backend_of(label):
    text = str(label).strip()
    for backend in KNOWN_BACKENDS:
        if text == backend or text.startswith(backend + "_"):
            return backend
    return None


def backends_of(node_labels):
    found = []
    for label in node_labels:
        if str(label).strip().endswith(BENCHMARK_SUFFIX):
            continue
        backend = backend_of(label)
        if backend and backend not in found:
            found.append(backend)
    return found


def read_sheet(workbook, name):
    if name not in workbook.sheetnames:
        return [], []
    rows = list(workbook[name].iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    return header, rows[1:]


def cell(row, index):
    return row[index] if index is not None and index < len(row) else None


def blank(value):
    if value is None:
        return True
    return str(value).strip().lower() in EMPTY_CELLS


def find_col(header, *candidates):
    for candidate in candidates:
        if candidate in header:
            return header.index(candidate)
    for candidate in candidates:
        for index, column in enumerate(header):
            if column.endswith("_" + candidate):
                return index
    return None


def node_cols(header, suffix):
    result = {}
    for index, column in enumerate(header):
        if column.endswith("_" + suffix):
            result[column[: -(len(suffix) + 1)]] = index
    return result


def nodes_from_header(header):
    found = set()
    for suffix in NODE_METRIC_SUFFIXES:
        found.update(node_cols(header, suffix))
    return {node for node in found if backend_of(node)}


def pick_data_col(header, rows, suffix):
    candidates = [
        index
        for index, column in enumerate(header)
        if column == suffix or column.endswith("_" + suffix)
    ]
    if not candidates:
        return None, None
    best = max(
        candidates,
        key=lambda index: sum(not blank(cell(row, index)) for row in rows),
    )
    return best, header[best]


def family_filled(header, rows, *suffixes):
    best = 0
    for index, column in enumerate(header):
        if any(column == suffix or column.endswith("_" + suffix) for suffix in suffixes):
            filled = sum(not blank(cell(row, index)) for row in rows)
            best = max(best, filled)
    return best


def detect_task(header, rows):
    accuracy = family_filled(header, rows, ACC_PASS)
    performance = family_filled(header, rows, PERF_TIME, PERF_PASS)
    if accuracy and performance:
        print(
            "精度列和性能列同时包含数据。按数据较多的任务解析。",
            file=sys.stderr,
        )
    if accuracy and accuracy >= performance:
        return "accuracy"
    if performance:
        return "performance"
    return None


def parse_outputs(detail):
    """提取精度详情中的逐输出结果。"""
    if not detail:
        return {}
    try:
        node_map = ast.literal_eval(str(detail))
        outputs = next(iter(node_map.values()))
        return {
            output["filename"]: bool(output.get("result"))
            for output in outputs
            if "filename" in output
        }
    except (SyntaxError, ValueError, TypeError, AttributeError, StopIteration):
        return {}


def to_float(value):
    """将单元格值转为四位小数。"""
    try:
        return round(float(str(value).strip()), 4)
    except (TypeError, ValueError):
        return None
